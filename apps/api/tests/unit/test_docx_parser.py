from docx import Document

from app.rag.ingestion.docx_parser import DocxParser


def test_docx_parser_extracts_paragraphs_and_tables_in_document_order(tmp_path) -> None:  # noqa: ANN001
    path = tmp_path / "handbook.docx"
    document = Document()
    document.add_paragraph("  First\tparagraph.  ")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Team"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "Support"
    table.cell(1, 1).text = "Ari"
    document.add_paragraph("Final paragraph.")
    document.save(path)

    blocks = DocxParser().parse(path)

    assert [block.metadata["block_type"] for block in blocks] == [
        "paragraph",
        "table",
        "paragraph",
    ]
    assert blocks[0].text == "First paragraph."
    assert blocks[1].text == "Team | Owner Support | Ari"
    assert blocks[2].text == "Final paragraph."
    assert all(block.page_number is None for block in blocks)
    assert all(block.metadata["source"] == "docx" for block in blocks)


def test_docx_parser_returns_no_blocks_when_document_has_no_text(tmp_path) -> None:  # noqa: ANN001
    path = tmp_path / "empty.docx"
    Document().save(path)

    assert DocxParser().parse(path) == []
